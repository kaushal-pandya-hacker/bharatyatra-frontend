import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# 1. Discover all Next.js pages in app/
app_dir = 'app'
pages = []
for root, dirs, files in os.walk(app_dir):
    for f in files:
        if f.startswith('page.'):
            full_path = os.path.join(root, f)
            rel_path = os.path.relpath(full_path, '.').replace(os.sep, '/')
            sub = os.path.relpath(root, app_dir).replace(os.sep, '/')
            if sub == '.':
                route_url = '/'
            else:
                parts = [p for p in sub.split('/') if not (p.startswith('(') and p.endswith(')'))]
                route_url = '/' + '/'.join(parts) if parts else '/'
            pages.append((route_url, rel_path))

pages.sort(key=lambda x: x[0])

# Page metadata catalog
PAGE_METADATA = {
    # Home & Global
    '/': ('BharatYatra Sovereign Homepage & Hero Portal', 'Public', 'Main marketing hero, AI travel prompt bar, popular circuits, and state guides'),
    '/feedback': ('User Feedback & Platform Experience Survey', 'Public / Traveler', 'Platform satisfaction rating, review form, and feedback collection'),
    '/privacy': ('Privacy Policy & Data Sovereign Governance', 'Public', 'Data protection terms, cookie policies, and GDPR / DPDP compliance'),
    '/terms': ('Terms of Service & Traveler Agreement', 'Public', 'Legal contract, cancellations, refunds, liability, and operator rules'),

    # Admin Portal (app/admin)
    '/admin': ('Master Admin Platform Executive Dashboard', 'Super Admin / Platform Ops', 'Real-time booking volume, revenue metrics, and operator health telemetry'),
    '/admin/audit-logs': ('System Security & Access Audit Trails', 'Super Admin / Platform Ops', 'Immutable ledger of administrative actions, permission escalations, and changes'),
    '/admin/bookings': ('Global Booking Oversight & Inventory Matrix', 'Super Admin / Platform Ops', 'All platform bookings across buses, trains, stays, and tour circuits'),
    '/admin/communications': ('Notification Dispatch & Campaign Center', 'Super Admin / Platform Ops', 'Bulk WhatsApp, SMS, and email dispatches to travelers and operators'),
    '/admin/inventory': ('Fleet, Room & Ticket Inventory Management', 'Super Admin / Platform Ops', 'Real-time allocation and locks on buses, hotels, and tourist vehicles'),
    '/admin/login': ('Admin Security Portal Authentication', 'Public (Staff)', 'Secure MFA login for platform engineers and system administrators'),
    '/admin/packages': ('Platform Package Oversight & Publishing', 'Super Admin / Platform Ops', 'Catalog management for national circuits, pricing tiers, and active status'),
    '/admin/packages/images': ('Package Media Asset & Cloudinary Manager', 'Super Admin / Platform Ops', 'High-res image CDN uploads, gallery tagging, and thumbnail optimization'),
    '/admin/payments': ('Payment Transaction Ledger & Gateway Telemetry', 'Super Admin / Platform Ops', 'Live payment stream, captured orders, UPI intents, and refund reconciliation'),
    '/admin/refunds': ('Refund Approvals, Chargebacks & Disputes', 'Super Admin / Platform Ops', 'Manual review queue for cancellation refunds and chargeback arbitration'),
    '/admin/settlements': ('Partner & Supplier Payout Settlement Engine', 'Super Admin / Platform Ops', 'Net payout calculation, commission withholding, and bank transfers'),
    '/admin/suppliers': ('Supplier Ecosystem Directory & Management', 'Super Admin / Platform Ops', 'List of onboarded hotel owners, transport fleets, and tour guides'),
    '/admin/suppliers/[id]': ('Supplier Performance & Contract Dossier', 'Super Admin / Platform Ops', 'Individual supplier contract, KYC documents, inventory, and reviews'),
    '/admin/tourism': ('Tourism Board Liaison & Regulatory Reporting', 'Super Admin / Platform Ops', 'State tourism board compliance, permit allocations, and tax reporting'),
    '/admin/travel-commissions': ('Commission Rates & Tier Settings', 'Super Admin / Platform Ops', 'Configurable percentage rates for agencies, affiliates, and ground fleets'),
    '/admin/travel-packages': ('Travel Partner Package Approval Queue', 'Super Admin / Platform Ops', 'Submission review pipeline for packages submitted by travel agencies'),
    '/admin/travel-packages/[id]/review': ('Partner Package Verification & Audit', 'Super Admin / Platform Ops', 'Detailed day-by-day itinerary verification, safety, and price validation'),
    '/admin/travel-partner-audit': ('Partner Audit Logs & Compliance Verification', 'Super Admin / Platform Ops', 'Audit records of partner tier updates, payouts, and license renewals'),
    '/admin/travel-partners': ('Partner Directory, Verification & Tiers', 'Super Admin / Platform Ops', 'Gold, Silver, Platinum partner status, booking quotas, and ratings'),
    '/admin/travel-partners/[id]': ('Partner KYC, Banking & Operational Dossier', 'Super Admin / Platform Ops', 'GSTIN verification, bank account details, and compliance records'),
    '/admin/users': ('Customer & Traveler Account Directory', 'Super Admin / Platform Ops', 'User account status, authentication logs, profile data, and roles'),

    # BharatYatra Operations (app/bharatyatra-ops)
    '/bharatyatra-ops': ('BharatYatra Operations Command Hub', 'Ground Ops / Dispatch', 'Live field logistics, airport/station fleet pickups, and guide dispatch'),
    '/bharatyatra-ops/audit-logs': ('Ground Operations Activity & Event Log', 'Ground Ops / Dispatch', 'Real-time operational events, driver assignments, and vehicle telemetry'),
    '/bharatyatra-ops/bookings': ('Operations Booking Tracking & Verification', 'Ground Ops / Dispatch', 'Verification of traveler check-ins, transit tickets, and guide vouchers'),
    '/bharatyatra-ops/dashboard': ('Ground Operations KPI & Field Telemetry', 'Ground Ops / Dispatch', 'Active convoys, weather warnings, and roadside assistance alerts'),
    '/bharatyatra-ops/login': ('Operations Staff Authentication Portal', 'Public (Ops Staff)', 'Secured login for dispatch coordinators and ground fleet managers'),
    '/bharatyatra-ops/notifications': ('Operational Alerts & Incident Feeds', 'Ground Ops / Dispatch', 'SOS alerts, route diversions, roadblock notifications, and delay reports'),
    '/bharatyatra-ops/settings': ('Ground Operations Configuration & Rules', 'Ground Ops / Dispatch', 'Pacing guidelines, buffer parameters, emergency contacts, and shifts'),
    '/bharatyatra-ops/trips': ('Live Road Trips & Convoy Monitoring', 'Ground Ops / Dispatch', 'Fleet tracking, driver contact links, and real-time waypoint radar'),
    '/bharatyatra-ops/trips/[id]': ('Convoy Telemetry & Incident Command', 'Ground Ops / Dispatch', 'Individual trip GPS coordinates, passenger roster, and vehicle diagnostics'),
    '/bharatyatra-ops/users': ('Field Staff & Driver Directory', 'Ground Ops / Dispatch', 'Certified drivers, licensed regional guides, and local coordinators'),
    '/bharatyatra-ops/users/[id]': ('Field Staff Profile & Certification Dossier', 'Ground Ops / Dispatch', 'Driving licenses, police verification, and safety ratings'),

    # Authentication & Onboarding
    '/auth/callback': ('OAuth Callback & Session Verification', 'Public / Traveler', 'Supabase OAuth token exchange, secure cookie setting, and redirection'),
    '/auth/verify-otp': ('SMS / WhatsApp One-Time Password Verification', 'Public / Traveler', 'Mobile phone verification and session restoration'),
    '/login': ('Traveler Login & Authentication Portal', 'Public', 'Email, OTP, and social sign-in with persistent 30-day session security'),
    '/register': ('New Traveler Account Registration', 'Public', 'Traveler registration with preferences, home city, and identity info'),
    '/signup': ('Traveler Quick Signup & Welcome Flow', 'Public', 'One-click account setup with instant itinerary vault creation'),
    '/beta-onboarding': ('VIP Early Access & Beta Traveler Onboarding', 'Invited Travelers', 'Early access features, special coupons, and prototype feedback'),
    '/onboarding/travel-profile': ('Traveler Personality & Preference Setup', 'Authenticated Traveler', 'Pacing, food preferences, accessibility requirements, and budget choices'),

    # BharatGateway FinTech Suite
    '/payment-gateway': ('BharatGateway™ FinTech Architecture Overview', 'Public / Developers', 'Enterprise payment infrastructure, multi-bank routing, and UPI docs'),
    '/payment-gateway/admin': ('BharatGateway Merchant & Transaction Terminal', 'FinTech Super Admin', 'Double-entry ledger balances, payout settlement queues, and risk scores'),
    '/payment-gateway/checkout/[orderId]': ('Hosted Sovereign Checkout Experience', 'Public / Traveler', 'PCI-DSS compliant payment interface for UPI, Cards, NetBanking, and Wallets'),
    '/payment-gateway/docs': ('BharatGateway Developer API & SDK Docs', 'Public / Developers', 'Webhooks, idempotent transaction creation, and SDK integration guides'),
    '/payment-gateway/merchant': ('Merchant Portal & API Key Management', 'Registered Merchant', 'Publishable and secret API keys, webhook signing secrets, and reports'),
    '/payment-gateway/pricing': ('FinTech Transaction Fees & Merchant Tiers', 'Public / Merchants', 'MDR rates for UPI (0%), RuPay Debit, Credit Cards, and net settlement'),

    # Travel Partner Portal (app/partner & app/partners)
    '/partner/bookings': ('Partner Agency Booking Portfolio', 'Travel Partner', 'Bookings dispatched to this specific partner with status and guest roster'),
    '/partner/commissions': ('Partner Commission & Payout Ledger', 'Travel Partner', 'Earned commissions, payout dates, TDS deductions, and bank invoices'),
    '/partner/dashboard': ('Travel Partner Command Center', 'Travel Partner', 'Monthly sales, active tours, client satisfaction score, and tier progress'),
    '/partner/notifications': ('Partner Updates & Booking Alerts', 'Travel Partner', 'Instant booking notifications, cancellation alerts, and admin approvals'),
    '/partner/packages': ('Partner Handcrafted Package Catalog', 'Travel Partner', 'Custom packages created by this agency with pricing and itineraries'),
    '/partner/packages/[id]/edit': ('Package Itinerary & Pricing Editor', 'Travel Partner', 'Modify tour highlights, departure dates, hotel inclusions, and pricing'),
    '/partner/packages/new': ('Create New Handcrafted Tour Package', 'Travel Partner', 'Draft multi-day tour package with day-by-day itinerary and photos'),
    '/partner/profile': ('Partner Agency Business Profile & Branding', 'Travel Partner', 'Agency logo, description, registered address, and customer ratings'),
    '/partner/settings': ('Partner Account & Banking Configuration', 'Travel Partner', 'Bank account details, GST number, payout preferences, and staff logins'),
    '/partners/application-status': ('Partner Application Tracking Portal', 'Applicant Agency', 'Track registration review status and licensing approval by BharatYatra'),
    '/partners/register': ('Travel Agency & Tour Operator Partnership Application', 'Public Agencies', 'Onboarding form for travel agents, ground handlers, and fleet operators'),

    # Supplier Portal (app/supplier)
    '/supplier': ('Supplier Platform Overview & Dashboard', 'Registered Supplier', 'Overview of properties, vehicle fleets, and daily booking requests'),
    '/supplier/activities': ('Activity & Excursion Inventory Management', 'Registered Supplier', 'Manage safari slots, heritage walks, river rafting, and VIP tickets'),
    '/supplier/bookings': ('Supplier Booking Queue & Confirmations', 'Registered Supplier', 'Accept or reject incoming booking requests and view voucher details'),
    '/supplier/bookings/[id]': ('Supplier Booking Voucher & Manifest', 'Registered Supplier', 'Guest contact, check-in date, special requests, and payout amounts'),
    '/supplier/dashboard': ('Supplier Revenue & Occupancy Analytics', 'Registered Supplier', 'Weekly occupancy, vehicle utilization rates, and customer reviews'),
    '/supplier/hotels': ('Hotel & Resort Room Inventory Controller', 'Registered Supplier', 'Room categories, seasonal rates, meal plan options, and block dates'),
    '/supplier/inventory': ('Master Real-Time Inventory Matrix', 'Registered Supplier', 'Live availability calendar across all supplier accommodations and transit'),
    '/supplier/login': ('Supplier Portal Authentication', 'Public (Suppliers)', 'Supplier account access with phone OTP and email credentials'),
    '/supplier/payments': ('Supplier Invoices & Remittance History', 'Registered Supplier', 'Cleared payouts, pending settlements, and tax withholding certificates'),
    '/supplier/profile': ('Supplier Business & Facility Profile', 'Registered Supplier', 'Hotel amenities, fleet specs, driver languages, and property photos'),
    '/supplier/register': ('New Supplier & Property Registration', 'Public Suppliers', 'Register hotel, homestay, tempo traveller, or tourist transport fleet'),
    '/supplier/restaurants': ('Dining & Culinary Venue Management', 'Registered Supplier', 'Restaurant table reservations, thali menus, and dining packages'),

    # Bookings, Cart & Checkout
    '/booking/confirmation': ('Booking Confirmation & Printable Receipt', 'Authenticated Traveler', 'Verified booking ID, downloadable invoice, and WhatsApp summary'),
    '/booking/failed': ('Payment Failure Recovery & Assistance', 'Traveler', 'Failed transaction details, retry with alternate UPI/Card, and support link'),
    '/booking/redirect': ('Gateway 3DS & Secure Bank Redirection', 'Traveler', 'Intermediate redirection bridge during banking authentication'),
    '/booking/success': ('Booking Success & Itinerary Dispatch', 'Authenticated Traveler', 'Immediate access to hotel vouchers, bus tickets, and trip command center'),
    '/bookings': ('My Bookings & Travel History', 'Authenticated Traveler', 'Complete list of past and upcoming hotel, bus, flight, and tour bookings'),
    '/bookings/[id]': ('Booking Detail, Voucher & QR Pass', 'Authenticated Traveler', 'Detailed ticket voucher, printable QR check-in code, and cancellation option'),
    '/bookings/[id]/pay': ('Pending Booking Checkout & Payment', 'Authenticated Traveler', 'Complete payment for reserved or partially paid itineraries'),
    '/cart': ('Travel Cart & Multi-Product Checkout Drawer', 'Authenticated Traveler', 'Bundled hotels, buses, tours, and activities ready for checkout'),
    '/checkout': ('Unified Sovereign Checkout Terminal', 'Authenticated Traveler', 'Traveler details, GST invoice option, coupon codes, and payment options'),
    '/payment': ('Generic Payment Bridge Terminal', 'Traveler', 'Fallback payment handler for external bookings and custom disbursements'),

    # Trip Command Center & Splitter
    '/trips': ('My Trips & Journey Vault (Command Hub)', 'Authenticated Traveler', 'Elevated trip dashboard with logo, travel telemetry, radar, and curated circuits'),
    '/trips/[id]': ('Trip Command Center & Live Telemetry', 'Authenticated Traveler', 'Real-time waypoint radar, vehicle diagnostics, itinerary nodes, and emergency SOS'),
    '/trips/[id]/cancel/[bookingId]': ('Trip Item Cancellation & Refund Calculator', 'Authenticated Traveler', 'Cancellation policy calculation and refund breakdown to original payment method'),
    '/trips/[id]/cost': ('Trip Cost Intelligence & Predictive Budget', 'Authenticated Traveler', 'Category-wise expense breakdown, fuel cost predictions, and buffer safety'),
    '/trips/[id]/expense': ('Trip Expense Quick Logger', 'Authenticated Traveler', 'Log new fuel, toll, food, or stay expense with receipt and payer tags'),
    '/trips/[id]/expenses': ('Group Expense Split & Multi-Payer Wallet', 'Authenticated Traveler', 'Multi-payer expense ledger with minor-unit mathematical precision and balances'),
    '/trips/[id]/settle': ('Debt Settlement & Payment Reconciliation', 'Authenticated Traveler', 'One-click UPI payment between group members to settle outstanding shares'),
    '/trips/[id]/settlement': ('Final Circuit Settlement & Archive', 'Authenticated Traveler', 'Finalized group balance statement and audit record upon trip completion'),
    '/trips/road-trip': ('Interactive Road Trip Navigator & Fuel Planner', 'Traveler', 'Live mileage calculation, toll estimates, and verified highway dhabas'),

    # Customer Discovery & Travel Products
    '/activities/[id]': ('Activity Details & Direct Booking', 'Public / Traveler', 'Tour details, timing slots, inclusions, guide info, and reviews'),
    '/ai-planner': ('Nexi AI Custom Itinerary Generator', 'Public / Traveler', 'Generate custom multi-day plans based on origin, duration, and pace'),
    '/buses': ('Intercity Bus Discovery & Route Finder', 'Public / Traveler', 'GSRTC, Volvo, and luxury sleeper buses with timing and fare filters'),
    '/buses/[id]': ('Bus Seat Selection & Boarding Point Selector', 'Public / Traveler', 'Interactive 2+1 and 2+2 seat layout, pickup points, and amenities'),
    '/circuits': ('Signature Pan-Indian Travel Circuits', 'Public / Traveler', 'Handcrafted regional circuits like Royal Kutch, Coastal Gujarat, and Golden Triangle'),
    '/circuits/[id]': ('Circuit Day-by-Day Master Itinerary', 'Public / Traveler', 'Interactive map route, daily schedule, included stays, and pricing'),
    '/dashboard': ('Traveler Personal Dashboard & Travel Vault', 'Authenticated Traveler', 'Saved trips, quick booking links, loyalty points, and recent searches'),
    '/destinations': ('Pan-India Destination Directory (36 States)', 'Public / Traveler', 'Explore top cities, spiritual anchors, hill stations, and heritage sites'),
    '/destinations/[id]': ('Destination Comprehensive Travel Guide', 'Public / Traveler', 'Weather, best time to visit, top sights, food recommendations, and stays'),
    '/expenses': ('Global Travel Expenses Overview', 'Authenticated Traveler', 'Cross-trip expense summaries, annual travel spend, and saved receipts'),
    '/experiences': ('Curated Cultural & Adventure Experiences', 'Public / Traveler', 'Local craft workshops, desert safaris, culinary trails, and village tours'),
    '/experiences/[id]': ('Experience Booking & Ticket Pass', 'Public / Traveler', 'Host profile, experience duration, meeting spot, and ticket availability'),
    '/explore': ('Interactive Indian Travel Explorer & AI Recommendations', 'Public / Traveler', 'Discovery feed tailored by budget, theme (Spiritual, Adventure, Heritage), and season'),
    '/flights': ('Domestic Flights Discovery & Fare Comparison', 'Public / Traveler', 'Direct flights, budget carriers, baggage allowances, and departure times'),
    '/flights/[id]': ('Flight Booking & Passenger Manifest', 'Public / Traveler', 'Seat selection, meal add-ons, passenger details, and airline baggage rules'),
    '/map': ('Interactive GIS Map Discovery & Route Visualizer', 'Public / Traveler', 'Full-screen Leaflet map displaying landmarks, stops, and highway routes'),
    '/my-trips': ('Saved Trips Alias & Vault Redirect', 'Authenticated Traveler', 'Alias route seamlessly linked to the primary /trips command center'),
    '/packages': ('Curated Tour Packages Catalog', 'Public / Traveler', 'Family tours, honeymoon packages, weekend getaways, and spiritual yatras'),
    '/packages/[slug]': ('Tour Package Detail, Dates & Pricing Calculator', 'Public / Traveler', 'Tiered pricing (Budget, Premium, Luxury), date calendar, and inclusions'),
    '/plan': ('Interactive AI Trip Planner & Route Builder', 'Public / Traveler', 'Personalized day-by-day planner with custom destinations and instant save'),
    '/profile': ('Traveler Profile, Preferences & Security Settings', 'Authenticated Traveler', 'Name, contact, travel ID proofs, emergency contact, and language'),
    '/profile/notifications': ('Traveler Notification Preferences & History', 'Authenticated Traveler', 'WhatsApp trip updates, booking alerts, and marketing preference toggles'),
    '/restaurants': ('Curated Dining & Regional Food Guide', 'Public / Traveler', 'Authentic Kathiyawadi, Gujarati thali, street food, and fine dining'),
    '/restaurants/[id]': ('Restaurant Menu, Hours & Table Booking', 'Public / Traveler', 'Signature dishes, pricing for two, hygiene ratings, and table reservations'),
    '/saved': ('Wishlist & Saved Places Collection', 'Authenticated Traveler', 'Bookmarked hotels, itineraries, packages, and hidden gems'),
    '/states': ('All 36 States & Union Territories of India', 'Public / Traveler', 'Complete geographic directory of Indian states with cultural overviews'),
    '/states/[slug]': ('State Tourism Portal & Regional Circuits', 'Public / Traveler', 'State highlights, top districts, regional cuisine, and tourism circuits'),
    '/states/loading': ('State Tourism Skeleton Preloader', 'System / Client', 'Optimized loading skeleton for fluid high-speed page transitions'),
    '/stays': ('Verified Accommodations, Heritage Haveli & Resorts', 'Public / Traveler', 'Heritage hotels, tented cities, homestays, and budget guest houses'),
    '/stays/[id]': ('Stay Details, Room Selection & Amenities', 'Public / Traveler', 'Photo gallery, room categories, breakfast options, and guest reviews'),
    '/trains': ('IRCTC Train Discovery & Live Schedule', 'Public / Traveler', 'Vande Bharat, Rajdhani, and express trains with route stops and coach classes'),
    '/trains/[id]': ('Train Route Details, Station Stops & Fare Matrix', 'Public / Traveler', 'Intermediate station halt timings, distance km, and class availability'),
}

# Group into clean domains
CATEGORIES = [
    ('Admin Portal (Super-Admin & Operations)', [p for p in pages if p[0].startswith('/admin')]),
    ('BharatYatra Ops Control Center', [p for p in pages if p[0].startswith('/bharatyatra-ops')]),
    ('Trip Command Center, Splitter & Telemetry', [p for p in pages if p[0].startswith('/trips')]),
    ('Bookings, Checkout & Settlement', [p for p in pages if p[0].startswith('/booking') or p[0] in ['/cart', '/checkout', '/payment']]),
    ('Authentication, Security & Onboarding', [p for p in pages if p[0].startswith('/auth') or p[0] in ['/login', '/register', '/signup', '/beta-onboarding'] or p[0].startswith('/onboarding')]),
    ('BharatGateway FinTech Suite', [p for p in pages if p[0].startswith('/payment-gateway')]),
    ('Travel Partner Ecosystem', [p for p in pages if p[0].startswith('/partner') or p[0].startswith('/partners')]),
    ('Supplier & Vendor Portal', [p for p in pages if p[0].startswith('/supplier')]),
    ('Customer Discovery, Itinerary & Travel Products', [
        p for p in pages if not (
            p[0].startswith('/admin') or
            p[0].startswith('/bharatyatra-ops') or
            p[0].startswith('/trips') or
            p[0].startswith('/booking') or
            p[0] in ['/cart', '/checkout', '/payment'] or
            p[0].startswith('/auth') or
            p[0] in ['/login', '/register', '/signup', '/beta-onboarding'] or
            p[0].startswith('/onboarding') or
            p[0].startswith('/payment-gateway') or
            p[0].startswith('/partner') or
            p[0].startswith('/partners') or
            p[0].startswith('/supplier')
        )
    ]),
]

# Verify all pages are covered
total_categorized = sum(len(cat[1]) for cat in CATEGORIES)
assert total_categorized == len(pages), f'Categorization mismatch: {total_categorized} vs {len(pages)}'

print(f'Total pages to document: {len(pages)}')
for cat_name, cat_pages in CATEGORIES:
    print(f'  - {cat_name}: {len(cat_pages)} pages')

# -------------------------------------------------------------
# BUILD MICROSOFT WORD (.DOCX) DOCUMENT
# -------------------------------------------------------------
doc = docx.Document()

# Page Margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

# Styling Helper
def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

# Title
p_title = doc.add_paragraph()
r_title = p_title.add_run('BHARAT YATRA')
r_title.font.size = Pt(24)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(11, 28, 48) # Navy #0b1c30

r_subtitle = p_title.add_run('\nComplete Pages & Routes Architectural Directory')
r_subtitle.font.size = Pt(14)
r_subtitle.font.bold = True
r_subtitle.font.color.rgb = RGBColor(230, 108, 0) # Amber / Saffron

p_meta = doc.add_paragraph()
p_meta.paragraph_format.space_after = Pt(14)
r_meta = p_meta.add_run('Production Platform: https://www.bharatyatra.in\nTotal Registered Pages: 126 Pages | Framework: Next.js 14 App Router\nCoverage: 100% Comprehensive Coverage across Traveler, Admin, Ops, Partner & FinTech Portals')
r_meta.font.size = Pt(9.5)
r_meta.font.color.rgb = RGBColor(100, 116, 139)

# Summary Table of Contents
doc.add_heading('Executive Domain Summary', level=2)
toc_table = doc.add_table(rows=1, cols=3)
toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
toc_table.autofit = False

hdr_cells = toc_table.rows[0].cells
hdr_cells[0].text = 'Application Subsystem'
hdr_cells[1].text = 'Registered Pages'
hdr_cells[2].text = 'Primary Scope'
for c in hdr_cells:
    set_cell_background(c, '0B1C30')
    set_cell_margins(c, top=120, bottom=120, left=150, right=150)
    for p in c.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5)

col_widths_toc = [Inches(3.2), Inches(1.3), Inches(2.5)]
for row in toc_table.rows:
    for idx, width in enumerate(col_widths_toc):
        row.cells[idx].width = width

for cat_name, cat_pages in CATEGORIES:
    row = toc_table.add_row()
    for idx, width in enumerate(col_widths_toc):
        row.cells[idx].width = width
    row.cells[0].text = cat_name
    row.cells[1].text = f'{len(cat_pages)} Pages'
    scope_text = 'Super Admin' if 'Admin' in cat_name else 'Ground Fleet & Ops' if 'Ops' in cat_name else 'FinTech Infrastructure' if 'BharatGateway' in cat_name else 'B2B Partners' if 'Partner' in cat_name else 'Hotels & Transit' if 'Supplier' in cat_name else 'Travelers & Public'
    row.cells[2].text = scope_text
    for cell in row.cells:
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)

doc.add_paragraph().paragraph_format.space_after = Pt(14)

# Detailed Tables by Category
for cat_name, cat_pages in CATEGORIES:
    h = doc.add_heading(f'{cat_name} ({len(cat_pages)} Pages)', level=1)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(2.4), Inches(1.8), Inches(1.8), Inches(1.0)]

    headers = ['Page Name & Purpose', 'Route URL', 'File Location', 'Access Level']
    for idx, text in enumerate(headers):
        c = table.rows[0].cells[idx]
        c.width = col_widths[idx]
        c.text = text
        set_cell_background(c, '0B1C30')
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(8.5)

    for i, (route, file_path) in enumerate(cat_pages):
        row = table.add_row()
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

        name, access, desc = PAGE_METADATA.get(route, (f'Page: {route}', 'Traveler', 'BharatYatra travel platform page'))

        row.cells[0].text = f'{name}\n{desc}'
        row.cells[1].text = route
        row.cells[2].text = file_path
        row.cells[3].text = access

        bg_color = 'F8FAFC' if i % 2 == 1 else 'FFFFFF'
        for idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
                    if idx == 0 and r.text.startswith(name):
                        r.font.bold = True
                    elif idx == 1:
                        r.font.color.rgb = RGBColor(180, 83, 9) # Amber
                    elif idx == 2:
                        r.font.color.rgb = RGBColor(71, 85, 105) # Slate

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

# Save .docx file
docx_path = 'BharatYatra_Pages_Directory.docx'
doc.save(docx_path)
print(f'Successfully generated {docx_path} ({os.path.getsize(docx_path)} bytes)')

# -------------------------------------------------------------
# BUILD .DOCS AND .MD FILES
# -------------------------------------------------------------
def build_markdown_document():
    lines = []
    lines.append('# BharatYatra — Complete Pages & Routes Directory\n')
    lines.append('> **Total Registered Pages:** **126**  ')
    lines.append('> **Framework:** Next.js 14 App Router  ')
    lines.append('> **Production URL:** https://www.bharatyatra.in  ')
    lines.append('> **Last Updated:** October 2026  ')
    lines.append('> **Coverage:** 100% of all App Router routes across Traveler Experience, Admin, Ops, Partner Ecosystem, FinTech BharatGateway, and Supplier Portals.\n')
    lines.append('## Table of Contents\n')

    for cat_name, cat_pages in CATEGORIES:
        anchor = cat_name.lower().replace(' ', '-').replace('&', '').replace(',', '').replace('(', '').replace(')', '').replace('--', '-')
        lines.append(f'- [{cat_name} ({len(cat_pages)} pages)](#{anchor})')
    lines.append('\n---\n')

    for cat_name, cat_pages in CATEGORIES:
        lines.append(f'## {cat_name}\n')
        lines.append('| Page Name & Purpose | Route URL | File Location | Access Level | Description |')
        lines.append('| :--- | :--- | :--- | :--- | :--- |')
        for route, file_path in cat_pages:
            name, access, desc = PAGE_METADATA.get(route, (f'Page: {route}', 'Traveler', 'BharatYatra platform page'))
            lines.append(f'| **{name}** | `{route}` | [`{file_path}`](file:///{os.path.abspath(file_path).replace(os.sep, "/")}) | `{access}` | {desc} |')
        lines.append('\n')

    return '\n'.join(lines)

md_content = build_markdown_document()

# Write to APP_PAGES_DIRECTORY.docs
docs_path = 'APP_PAGES_DIRECTORY.docs'
with open(docs_path, 'w', encoding='utf-8') as f:
    f.write(md_content)
print(f'Successfully generated {docs_path} ({os.path.getsize(docs_path)} bytes)')

# Update APP_PAGES_DIRECTORY.md
md_path = 'APP_PAGES_DIRECTORY.md'
with open(md_path, 'w', encoding='utf-8') as f:
    f.write(md_content)
print(f'Successfully updated {md_path} ({os.path.getsize(md_path)} bytes)')
